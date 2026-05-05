"""Unit tests for the docx / xlsx text scanner.

Fixtures are built in-memory so the test suite stays self-contained — no
binary files in the repo. This also makes failure modes easier to debug:
a broken assertion points at the constructor that produced the bytes.
"""

from __future__ import annotations

import io

import openpyxl
import pytest
from docx import Document
from docx.shared import Pt

from src.skills.template_filler.text_scanner import (
    TextFragment,
    chunk_for_prompt,
    scan_docx,
    scan_xlsx,
)


def _make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    """Build a minimal .docx with the given paragraphs and optional table."""
    doc = Document()
    for p in paragraphs:
        para = doc.add_paragraph(p)
        # Touch a property so python-docx serialises a non-empty run set
        if para.runs:
            para.runs[0].font.size = Pt(11)
    if table:
        rows = len(table)
        cols = max(len(r) for r in table) if table else 0
        t = doc.add_table(rows=rows, cols=cols)
        for ri, row in enumerate(table):
            for ci, val in enumerate(row):
                t.cell(ri, ci).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_xlsx(rows_per_sheet: dict[str, list[list]]) -> bytes:
    """Build a minimal .xlsx with multiple sheets."""
    wb = openpyxl.Workbook()
    # Workbook starts with one default sheet — repurpose it for the first
    # caller-supplied sheet to avoid orphan blanks.
    sheet_names = list(rows_per_sheet.keys())
    if not sheet_names:
        wb.remove(wb.active)
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()
    wb.active.title = sheet_names[0]
    for ri, row in enumerate(rows_per_sheet[sheet_names[0]], start=1):
        for ci, val in enumerate(row, start=1):
            wb.active.cell(row=ri, column=ci, value=val)
    for name in sheet_names[1:]:
        ws = wb.create_sheet(name)
        for ri, row in enumerate(rows_per_sheet[name], start=1):
            for ci, val in enumerate(row, start=1):
                ws.cell(row=ri, column=ci, value=val)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_scan_docx_paragraphs_in_order():
    raw = _make_docx(["客户:[客户名]", "", "金额:[金额]"])
    out = scan_docx(raw)

    # Empty paragraphs are dropped.
    assert [f.text for f in out] == ["客户:[客户名]", "金额:[金额]"]
    assert all(f.location_hint.startswith("段落 ") for f in out)


def test_scan_docx_includes_table_cells():
    raw = _make_docx(
        ["头部说明"],
        table=[["项目", "数量", "金额"], ["商品 [SKU]", "[数量]", "[金额]"]],
    )
    out = scan_docx(raw)
    texts = [f.text for f in out]

    assert "头部说明" in texts
    assert "商品 [SKU]" in texts
    assert "[数量]" in texts
    # Table cells get a "表 X · 行Y列Z" hint.
    assert any("表 1 ·" in f.location_hint for f in out)


def test_scan_xlsx_walks_all_sheets_in_order():
    raw = _make_xlsx(
        {
            "采购单": [["客户", "[客户名]"], ["日期", "[日期]"]],
            "明细": [["SKU", "数量"], ["[SKU]", 10]],
        }
    )
    out = scan_xlsx(raw)
    # Every visible cell should appear, with sheet!cell coordinate hint.
    by_hint = {f.location_hint: f.text for f in out}
    assert by_hint["采购单!A1"] == "客户"
    assert by_hint["采购单!B1"] == "[客户名]"
    assert by_hint["采购单!B2"] == "[日期]"
    assert by_hint["明细!A1"] == "SKU"
    assert by_hint["明细!A2"] == "[SKU]"
    # Non-string values get stringified.
    assert by_hint["明细!B2"] == "10"


def test_scan_xlsx_skips_blank_cells():
    raw = _make_xlsx({"S": [["a", None, "b"], [None, None, None], ["c"]]})
    out = scan_xlsx(raw)
    cell_texts = [f.text for f in out if not f.location_hint.endswith("!__overview__")]
    assert sorted(cell_texts) == ["a", "b", "c"]


def test_scan_xlsx_emits_overview_per_sheet():
    raw = _make_xlsx({"S1": [["a", "b"]], "S2": [["c"]]})
    out = scan_xlsx(raw)
    overviews = [f for f in out if f.location_hint.endswith("!__overview__")]
    assert {f.location_hint for f in overviews} == {"S1!__overview__", "S2!__overview__"}
    # Overview comes before the cells of its sheet so the LLM reads it first.
    s1_indices = [i for i, f in enumerate(out) if f.location_hint.startswith("S1!")]
    assert out[s1_indices[0]].location_hint == "S1!__overview__"


def test_scan_xlsx_overview_flags_single_header_row_form():
    """Header row + nothing else should still trigger the structural hint.

    Real-world entry forms often look exactly like this: one row of column
    labels at A1:H1 and openpyxl's max_row reports 1 (not "1 + blank pad").
    Without the explicit hint the LLM can't tell apart a tiny one-row note
    from a blank-row template, so Mode B never fires.
    """
    raw = _make_xlsx({"Sheet1": [["姓名", "工号", "部门"]]})
    out = scan_xlsx(raw)
    overview = next(f for f in out if f.location_hint == "Sheet1!__overview__")
    assert "结构提示" in overview.text
    assert "第 2 行" in overview.text


def test_scan_xlsx_overview_flags_header_with_blank_data_rows():
    """Header in row 1, max_row pushed to 5 by an explicit blank cell, no other data."""
    wb = openpyxl.Workbook()
    wb.active.title = "Sheet1"
    wb.active["A1"] = "姓名"
    wb.active["B1"] = "工号"
    # Touch row 5 to bump max_row past the header without putting a value
    # in any earlier blank row.
    wb.active.cell(row=5, column=1).value = None
    wb.active.cell(row=5, column=1).comment = None
    # openpyxl tracks max_row via writes — set then clear by writing empty.
    wb.active["A5"] = "x"
    wb.active["A5"] = None
    buf = io.BytesIO()
    wb.save(buf)
    out = scan_xlsx(buf.getvalue())
    overview = next(f for f in out if f.location_hint == "Sheet1!__overview__")
    assert "结构提示" in overview.text


def test_scan_xlsx_overview_no_hint_for_filled_table():
    """Multi-row sheets with data in every row aren't treated as forms."""
    raw = _make_xlsx(
        {
            "明细": [
                ["SKU", "数量"],
                ["A001", 10],
                ["A002", 20],
                ["A003", 30],
            ]
        }
    )
    out = scan_xlsx(raw)
    overview = next(f for f in out if f.location_hint == "明细!__overview__")
    assert "结构提示" not in overview.text


def test_chunk_for_prompt_respects_max_chars():
    fragments = [TextFragment(text="a" * 100, location_hint=f"loc{i}") for i in range(50)]
    batches = list(chunk_for_prompt(fragments, max_chars=400))

    # Each batch's payload must stay under the budget (allow some slack
    # for hint/overhead per fragment).
    for batch in batches:
        joined = sum(len(f.text) + len(f.location_hint) + 4 for f in batch)
        assert joined <= 500
    # No fragment was lost.
    assert sum(len(b) for b in batches) == 50


def test_chunk_for_prompt_handles_empty_input():
    assert list(chunk_for_prompt([], max_chars=100)) == []


def test_scan_docx_rejects_corrupt_bytes():
    """Calling scan_docx on garbage bytes should raise — caller catches it."""
    with pytest.raises(Exception):
        scan_docx(b"definitely not a docx")
