"""Smoke tests for the /api/template/extract_fields router.

These hit the FastAPI route via TestClient with a stub LLM patched in,
so no real model call is made. Goal: verify the route plumbing (file
upload → extractor → JSON shape), not the extraction quality (which
is covered in test_extractor / test_llm_prompt).
"""

from __future__ import annotations

import io
from typing import Any
from unittest.mock import patch

import openpyxl
import pytest
from docx import Document
from fastapi.testclient import TestClient


@pytest.fixture
def app_with_router():
    """Build a minimal FastAPI app that mounts only the templates router.

    Booting the full gateway app would require live config loading, OSS
    credentials, model factory wiring — none of which are necessary to
    exercise this single endpoint.
    """
    from fastapi import FastAPI

    from src.gateway.routers import templates as tpl_router

    app = FastAPI()
    app.include_router(tpl_router.router)
    return app


def _docx_bytes() -> bytes:
    doc = Document()
    doc.add_paragraph("客户:[客户名]")
    doc.add_paragraph("日期:[日期]")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _xlsx_bytes() -> bytes:
    wb = openpyxl.Workbook()
    wb.active["A1"] = "客户"
    wb.active["B1"] = "[客户名]"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_fields_returns_json_shape(app_with_router):
    """End-to-end: post a docx, get back the documented response shape."""

    def stub_llm_factory(_model_name: Any):
        def call(_system: str, _user: str) -> str:
            return '[{"name":"customer_name","label":"客户","type":"string","required":true,"original_text":"[客户名]"}]'

        return call

    with patch("src.gateway.routers.templates._build_llm_callable", stub_llm_factory):
        client = TestClient(app_with_router)
        resp = client.post(
            "/api/template/extract_fields",
            files={"file": ("po.docx", _docx_bytes(), "application/octet-stream")},
        )

    assert resp.status_code == 200, resp.text
    body = resp.json()
    assert "fields" in body and "warnings" in body and "source_text_preview" in body
    assert len(body["fields"]) == 1
    assert body["fields"][0]["name"] == "customer_name"
    assert body["fields"][0]["original_text"] == "[客户名]"


def test_extract_fields_empty_upload_rejected(app_with_router):
    client = TestClient(app_with_router)
    resp = client.post(
        "/api/template/extract_fields",
        files={"file": ("empty.docx", b"", "application/octet-stream")},
    )
    assert resp.status_code == 400
    assert "empty" in resp.json()["detail"].lower()


def test_extract_fields_xlsx_path(app_with_router):
    """Confirms the .xlsx extension dispatches to the xlsx scanner."""

    def stub_llm_factory(_model_name: Any):
        return lambda _s, _u: "[]"  # nothing recognized — fields=[]

    with patch("src.gateway.routers.templates._build_llm_callable", stub_llm_factory):
        client = TestClient(app_with_router)
        resp = client.post(
            "/api/template/extract_fields",
            files={"file": ("po.xlsx", _xlsx_bytes(), "application/octet-stream")},
        )

    assert resp.status_code == 200
    assert resp.json()["fields"] == []
    # The preview should still come through even with zero fields.
    assert resp.json()["source_text_preview"] is not None


def test_extract_fields_unsupported_extension_returns_warning(app_with_router):
    """A .pdf doesn't blow up — extractor returns a warning, route is 200."""

    with patch("src.gateway.routers.templates._build_llm_callable", lambda *_a: lambda _s, _u: "[]"):
        client = TestClient(app_with_router)
        resp = client.post(
            "/api/template/extract_fields",
            files={"file": ("not_a_template.pdf", b"hello", "application/pdf")},
        )

    assert resp.status_code == 200
    body = resp.json()
    assert body["fields"] == []
    assert any(w.get("code") == "unsupported_extension" for w in body["warnings"])


def test_extract_fields_oversized_rejected(app_with_router):
    """Test the size-cap path without actually allocating 50 MiB."""
    from src.gateway.routers import templates as tpl_router

    with patch.object(tpl_router, "MAX_TEMPLATE_BYTES", 16):
        client = TestClient(app_with_router)
        resp = client.post(
            "/api/template/extract_fields",
            files={"file": ("po.docx", b"x" * 64, "application/octet-stream")},
        )

    assert resp.status_code == 400
    assert "too large" in resp.json()["detail"]
