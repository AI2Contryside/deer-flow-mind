"""Tests for renderers — dispatch logic + the openpyxl xlsx renderer.

The xlsx renderer is openpyxl-based now (xltpl was dropped because its
``BookWriter`` internally calls ``xlrd.open_workbook``, which xlrd 2.0+
no longer supports for xlsx files). The docx render path still pulls in
``docxtpl`` and skips when the optional dep is missing.
"""

from __future__ import annotations

import io

import openpyxl
import pytest

from src.skills.template_filler.renderers import RendererError, render_for_extension
from src.skills.template_filler.renderers.xlsx import render as render_xlsx


def _make_xlsx_with_jinja_tags() -> bytes:
    """Produce a tiny xlsx whose cells already carry `{{ name }}` markers."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "客户:"
    ws["B1"] = "{{ customer_name }}"
    ws["A2"] = "金额:"
    ws["B2"] = "金额是 {{ amount }} 元"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_render_for_extension_rejects_unknown():
    with pytest.raises(ValueError, match="unsupported extension"):
        render_for_extension("template.pdf", b"", {})


def test_xlsx_render_substitutes_simple_tags():
    src = _make_xlsx_with_jinja_tags()
    rendered = render_xlsx(src, {"customer_name": "Acme Hardware", "amount": 12345})

    wb = openpyxl.load_workbook(io.BytesIO(rendered))
    ws = wb.active
    assert ws["B1"].value == "Acme Hardware"
    assert ws["B2"].value == "金额是 12345 元"


def test_xlsx_render_leaves_missing_keys_intact():
    """Missing keys leave `{{ x }}` literal so the gap is visible upstream
    rather than silently emptying a cell."""
    src = _make_xlsx_with_jinja_tags()
    rendered = render_xlsx(src, {"customer_name": "A"})

    wb = openpyxl.load_workbook(io.BytesIO(rendered))
    ws = wb.active
    assert ws["B1"].value == "A"
    assert "{{ amount }}" in ws["B2"].value


def test_xlsx_render_handles_no_tags():
    """Cells without tags must round-trip unchanged."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = "no template here"
    buf = io.BytesIO()
    wb.save(buf)

    rendered = render_xlsx(buf.getvalue(), {"customer_name": "X"})
    wb2 = openpyxl.load_workbook(io.BytesIO(rendered))
    assert wb2.active["A1"].value == "no template here"


def test_xlsx_render_substitutes_none_as_empty_string():
    """A None value renders to "" so a deliberately-blank field doesn't print "None".

    openpyxl coerces empty-string cells back to ``None`` on round-trip, so
    we accept either shape — the contract is just "no literal 'None'
    written into the cell".
    """
    src = _make_xlsx_with_jinja_tags()
    rendered = render_xlsx(src, {"customer_name": None, "amount": 0})
    wb = openpyxl.load_workbook(io.BytesIO(rendered))
    assert wb.active["B1"].value in (None, "")
    assert wb.active["B2"].value == "金额是 0 元"


def test_xlsx_render_via_dispatch_works_end_to_end():
    """render_for_extension routes ``.xlsx`` to the openpyxl path."""
    src = _make_xlsx_with_jinja_tags()
    result = render_for_extension(
        "template.xlsx", src, {"customer_name": "X", "amount": 1}
    )
    # Dispatch returns a RenderResult envelope; .content holds the bytes.
    wb = openpyxl.load_workbook(io.BytesIO(result.content))
    assert wb.active["B1"].value == "X"
    assert "spreadsheetml" in result.content_type


def test_xlsx_render_wraps_corrupt_input_as_RendererError():
    """Bad xlsx bytes surface as RendererError, not raw openpyxl exception."""
    with pytest.raises(RendererError, match="xlsx render failed"):
        render_xlsx(b"\x00not a real xlsx", {})


def test_render_docx_when_docxtpl_missing_raises_importerror():
    try:
        import docxtpl  # noqa: F401

        pytest.skip("docxtpl is installed; ImportError path doesn't apply")
    except ImportError:
        pass

    with pytest.raises(ImportError, match="docxtpl"):
        render_for_extension("anything.docx", b"\x00\x00", {})


def test_renderer_error_class_distinguishes_from_importerror():
    """Sanity: RendererError is its own exception type."""
    err = RendererError("kaboom")
    assert isinstance(err, RuntimeError)
    assert not isinstance(err, ImportError)
