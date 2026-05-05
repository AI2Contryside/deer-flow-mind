"""Tests for the openpyxl/python-docx jinja-ifier."""

from __future__ import annotations

import io

import openpyxl
import pytest
from docx import Document

from src.skills.template_filler.jinjaify import jinjaify
from src.skills.template_filler.types import ExtractedField


def _xlsx_with_headers() -> bytes:
    """Two-row sheet: header row + blank data row, as a typical entry form."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws["A1"] = "姓名"
    ws["B1"] = "工号"
    ws["C1"] = "部门"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _xlsx_with_inline_placeholder() -> bytes:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws["A1"] = "客户:[客户名]"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _docx_with_inline_placeholder() -> bytes:
    doc = Document()
    doc.add_paragraph("客户:[客户名]")
    doc.add_paragraph("日期:[日期]")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ----------------------------- xlsx header-cell ----------------------------


def test_xlsx_header_cell_writes_jinja_into_data_row():
    """LLM identifies "姓名"=A1 with cell_anchor=A2 → A2 gets `{{ name }}`,
    A1 stays untouched."""
    src = _xlsx_with_headers()
    fields = [
        ExtractedField(name="name", label="姓名", original_text="", cell_anchor="Sheet1!A2"),
        ExtractedField(name="emp_id", label="工号", original_text="", cell_anchor="Sheet1!B2"),
    ]
    result = jinjaify(src, "people.xlsx", fields)

    wb = openpyxl.load_workbook(io.BytesIO(result.content))
    ws = wb.active
    # Headers preserved
    assert ws["A1"].value == "姓名"
    assert ws["B1"].value == "工号"
    # Data row got jinja tags
    assert ws["A2"].value == "{{ name }}"
    assert ws["B2"].value == "{{ emp_id }}"
    assert result.applied == ["name", "emp_id"]


def test_xlsx_cell_anchor_without_sheet_prefix_falls_back_to_active():
    src = _xlsx_with_headers()
    fields = [ExtractedField(name="name", label="姓名", original_text="", cell_anchor="A2")]
    result = jinjaify(src, "people.xlsx", fields)

    wb = openpyxl.load_workbook(io.BytesIO(result.content))
    assert wb.active["A2"].value == "{{ name }}"


def test_xlsx_malformed_anchor_skipped():
    src = _xlsx_with_headers()
    fields = [ExtractedField(name="bad", label="坏的", original_text="", cell_anchor="not-a-cell")]
    result = jinjaify(src, "people.xlsx", fields)
    assert result.applied == []
    assert result.skipped[0].reason == "malformed_cell_anchor"


def test_xlsx_unknown_sheet_skipped():
    src = _xlsx_with_headers()
    fields = [
        ExtractedField(name="x", label="X", original_text="", cell_anchor="DoesNotExist!A1"),
    ]
    result = jinjaify(src, "people.xlsx", fields)
    assert result.skipped[0].reason == "sheet_not_found"


# ----------------------------- xlsx inline-text ----------------------------


def test_xlsx_inline_text_replace():
    src = _xlsx_with_inline_placeholder()
    fields = [ExtractedField(name="customer_name", label="客户", original_text="[客户名]")]
    result = jinjaify(src, "form.xlsx", fields)

    wb = openpyxl.load_workbook(io.BytesIO(result.content))
    assert wb.active["A1"].value == "客户:{{ customer_name }}"
    assert result.applied == ["customer_name"]


def test_xlsx_inline_text_ambiguous_skipped():
    """Same placeholder in two cells → skipped."""
    wb = openpyxl.Workbook()
    wb.active["A1"] = "[名]"
    wb.active["B1"] = "[名]"
    buf = io.BytesIO()
    wb.save(buf)

    fields = [ExtractedField(name="name", label="名", original_text="[名]")]
    result = jinjaify(buf.getvalue(), "form.xlsx", fields)
    assert result.skipped[0].reason == "ambiguous"


def test_xlsx_inline_text_not_found():
    src = _xlsx_with_inline_placeholder()
    fields = [ExtractedField(name="missing", label="x", original_text="[xxx]")]
    result = jinjaify(src, "form.xlsx", fields)
    assert result.skipped[0].reason == "not_found"


# ----------------------------- docx ---------------------------------------


def test_docx_inline_text_replace():
    src = _docx_with_inline_placeholder()
    fields = [
        ExtractedField(name="customer_name", label="客户", original_text="[客户名]"),
        ExtractedField(name="date", label="日期", original_text="[日期]"),
    ]
    result = jinjaify(src, "po.docx", fields)

    doc = Document(io.BytesIO(result.content))
    paragraphs = [p.text for p in doc.paragraphs]
    assert any("{{ customer_name }}" in t for t in paragraphs)
    assert any("{{ date }}" in t for t in paragraphs)
    assert sorted(result.applied) == ["customer_name", "date"]


def test_docx_cell_anchor_field_unsupported():
    src = _docx_with_inline_placeholder()
    fields = [ExtractedField(name="x", label="X", original_text="", cell_anchor="Sheet1!A2")]
    result = jinjaify(src, "po.docx", fields)
    assert result.skipped[0].reason == "cell_anchor_unsupported_for_docx"


# ----------------------------- validation --------------------------------


def test_invalid_variable_name_skipped():
    src = _xlsx_with_headers()
    # digit prefix + reserved keyword
    fields = [
        ExtractedField(name="9bad", label="坏", original_text="", cell_anchor="Sheet1!A2"),
        ExtractedField(name="for", label="保留", original_text="", cell_anchor="Sheet1!B2"),
    ]
    result = jinjaify(src, "x.xlsx", fields)
    reasons = sorted(o.reason for o in result.skipped)
    assert reasons == ["invalid_variable_name", "invalid_variable_name"]


def test_unsupported_extension():
    with pytest.raises(ValueError, match="unsupported template extension"):
        jinjaify(b"raw", "x.pdf", [])


def test_field_with_no_target_skipped():
    src = _xlsx_with_headers()
    fields = [ExtractedField(name="x", label="X", original_text="", cell_anchor=None)]
    result = jinjaify(src, "x.xlsx", fields)
    assert result.skipped[0].reason == "no_target"
