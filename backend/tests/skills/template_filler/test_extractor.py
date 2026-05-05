"""End-to-end extractor tests with a deterministic stub LLM.

Real LLM behaviour is covered separately at integration time. Here we
verify the orchestration: scan → chunk → call → parse → dedupe → result.
"""

from __future__ import annotations

import io
import json

import openpyxl
from docx import Document

from src.skills.template_filler.extractor import extract_fields


def _make_simple_docx() -> bytes:
    doc = Document()
    doc.add_paragraph("客户:[客户名]")
    doc.add_paragraph("日期:[日期]")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _make_simple_xlsx() -> bytes:
    wb = openpyxl.Workbook()
    wb.active.title = "S"
    wb.active["A1"] = "客户"
    wb.active["B1"] = "[客户名]"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def test_extract_unsupported_extension_returns_warning():
    result = extract_fields(b"raw text", "template.txt")
    assert result.fields == []
    assert any(w.code == "unsupported_extension" for w in result.warnings)


def test_extract_empty_template_returns_warning():
    # Doc with only blank paragraphs scans to zero fragments.
    doc = Document()
    doc.add_paragraph("")
    buf = io.BytesIO()
    doc.save(buf)

    result = extract_fields(buf.getvalue(), "blank.docx")
    assert result.fields == []
    assert any(w.code == "no_text_extracted" for w in result.warnings)


def test_extract_with_stub_llm_returns_fields():
    """LLM stub returns a fixed JSON array — extractor folds it into result."""

    def stub_llm(_system: str, _user: str) -> str:
        return json.dumps(
            [
                {
                    "name": "customer_name",
                    "label": "客户",
                    "type": "string",
                    "required": True,
                    "original_text": "[客户名]",
                    "location_hint": "段落 1",
                    "description": "采购方公司",
                }
            ]
        )

    result = extract_fields(_make_simple_docx(), "po.docx", llm_callable=stub_llm)

    assert len(result.fields) == 1
    assert result.fields[0].name == "customer_name"
    assert result.fields[0].original_text == "[客户名]"
    assert result.warnings == []
    # Preview should contain at least the first paragraph's text.
    assert result.source_text_preview is not None
    assert "[客户名]" in result.source_text_preview


def test_extract_dedupe_across_chunks():
    """If the LLM returns the same field twice, dedupe keeps the first."""

    seen_calls = {"n": 0}

    def stub_llm(_system: str, _user: str) -> str:
        seen_calls["n"] += 1
        return json.dumps([{"name": "customer_name", "label": "客户", "original_text": "[客户名]"}])

    result = extract_fields(_make_simple_docx(), "po.docx", llm_callable=stub_llm)
    # Even if multiple chunks each return the same field, output is unique.
    assert [f.name for f in result.fields] == ["customer_name"]


def test_extract_records_llm_call_failure_as_warning():
    """LLM exception turns into a warning, not a hard failure."""

    def boom(*_args: object, **_kw: object) -> str:
        raise RuntimeError("model unavailable")

    result = extract_fields(_make_simple_docx(), "po.docx", llm_callable=boom)
    assert result.fields == []
    assert any(w.code == "llm_call_failed" for w in result.warnings)


def test_extract_records_unparseable_llm_reply():
    def stub(*_args: object, **_kw: object) -> str:
        return "我无法识别"

    result = extract_fields(_make_simple_docx(), "po.docx", llm_callable=stub)
    assert result.fields == []
    assert any(w.code == "llm_unparseable" for w in result.warnings)


def test_extract_xlsx_works_via_same_pipeline():
    """Smoke test that the xlsx path doesn't blow up downstream."""

    def stub_llm(_s: str, _u: str) -> str:
        return json.dumps([{"name": "customer_name", "label": "客户", "original_text": "[客户名]"}])

    result = extract_fields(_make_simple_xlsx(), "po.xlsx", llm_callable=stub_llm)
    assert len(result.fields) == 1
    assert result.fields[0].name == "customer_name"


def test_extract_corrupt_bytes_yields_scan_failed_warning():
    result = extract_fields(b"not a real docx", "broken.docx")
    assert result.fields == []
    assert any(w.code == "scan_failed" for w in result.warnings)


def test_extract_keeps_multiple_mode_b_fields():
    """Header-cell fields all share original_text="" — dedupe must not collapse them.

    Regression test for a bug where _dedupe_fields kept only the first
    cell-anchor field because it keyed on original_text and "" matched
    every subsequent one.
    """

    def stub_llm(_s: str, _u: str) -> str:
        return json.dumps(
            [
                {
                    "name": "name",
                    "label": "姓名",
                    "original_text": "",
                    "cell_anchor": "Sheet1!A2",
                },
                {
                    "name": "emp_id",
                    "label": "工号",
                    "original_text": "",
                    "cell_anchor": "Sheet1!B2",
                },
                {
                    "name": "department",
                    "label": "部门",
                    "original_text": "",
                    "cell_anchor": "Sheet1!C2",
                },
            ]
        )

    result = extract_fields(_make_simple_xlsx(), "people.xlsx", llm_callable=stub_llm)

    assert [f.name for f in result.fields] == ["name", "emp_id", "department"]
    assert all(f.cell_anchor for f in result.fields)
    assert all(f.original_text == "" for f in result.fields)


def test_extract_dedupes_same_cell_anchor():
    """Same cell_anchor on two fields → keep first only."""

    def stub_llm(_s: str, _u: str) -> str:
        return json.dumps(
            [
                {"name": "a", "label": "A", "original_text": "", "cell_anchor": "Sheet1!A2"},
                {"name": "a_dupe", "label": "A2", "original_text": "", "cell_anchor": "Sheet1!A2"},
                {"name": "b", "label": "B", "original_text": "", "cell_anchor": "Sheet1!B2"},
            ]
        )

    result = extract_fields(_make_simple_xlsx(), "x.xlsx", llm_callable=stub_llm)
    assert [f.name for f in result.fields] == ["a", "b"]
