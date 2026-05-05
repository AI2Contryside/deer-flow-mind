"""Pull visible text out of a docx / xlsx template, with position hints.

This is the deterministic half of field extraction — no LLM, no heuristics
about *which* text spans are placeholders. We just emit "here is the text,
here is where it lives" so the LLM gets a structured view it can reason
about with location hints intact.
"""

from __future__ import annotations

import io
from collections.abc import Iterator
from dataclasses import dataclass
from typing import IO

import openpyxl


@dataclass(frozen=True)
class TextFragment:
    """One chunk of visible text plus where it came from.

    `location_hint` is human-readable — used by the LLM prompt and surfaced
    on the review UI, never as a programmatic addressing scheme. The actual
    rewrite path uses `text` itself (matched by the Go jinja_renderer).
    """

    text: str
    location_hint: str


def scan_docx(file: bytes | IO[bytes]) -> list[TextFragment]:
    """Extract paragraph and table-cell text from a .docx file.

    `python-docx` already concatenates run text per paragraph, which is
    exactly what we want — placeholders split across runs come back fused.
    Empty paragraphs are dropped to keep the LLM context lean.
    """
    # Lazy import: docx is only needed for .docx templates. Importing at
    # module level would force every consumer (e.g. the xlsx-only path)
    # to install python-docx even when they never touch a .docx file.
    from docx import Document as DocxDocument

    if isinstance(file, bytes):
        file = io.BytesIO(file)
    doc = DocxDocument(file)
    out: list[TextFragment] = []

    for idx, para in enumerate(doc.paragraphs, start=1):
        text = para.text.strip()
        if not text:
            continue
        out.append(TextFragment(text=text, location_hint=f"段落 {idx}"))

    for table_idx, table in enumerate(doc.tables, start=1):
        for row_idx, row in enumerate(table.rows, start=1):
            for col_idx, cell in enumerate(row.cells, start=1):
                cell_text = cell.text.strip()
                if not cell_text:
                    continue
                out.append(
                    TextFragment(
                        text=cell_text,
                        location_hint=f"表 {table_idx} · 行{row_idx}列{col_idx}",
                    )
                )
    return out


def scan_xlsx(file: bytes | IO[bytes]) -> list[TextFragment]:
    """Extract every non-empty cell's display value from a .xlsx file.

    `data_only=False` is critical — we want the *formula or the literal*
    cell value, not a cached evaluated number. Sheets are scanned in
    workbook order; cells in row-major order within each sheet.

    Each sheet emits one ``<sheet>!__overview__`` fragment up front
    summarising row counts and the non-empty row numbers. That lets the
    LLM detect "header row + many blank data rows" templates — without
    the overview, the model only sees a handful of header cells in
    isolation and can't tell whether row 2+ is blank or just absent
    from the prompt.
    """
    if isinstance(file, bytes):
        file = io.BytesIO(file)
    wb = openpyxl.load_workbook(file, data_only=False, read_only=True)
    out: list[TextFragment] = []

    try:
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            cells: list[tuple[str, int, str]] = []  # (coord, row_idx, text)
            non_empty_rows: set[int] = set()
            max_row_seen = 0
            for row in ws.iter_rows():
                for cell in row:
                    # `EmptyCell` in read-only mode is a placeholder with no
                    # row/coordinate attrs — getattr lets us count blanks
                    # for max_row tracking without crashing.
                    cell_row = getattr(cell, "row", None)
                    if cell_row is not None:
                        max_row_seen = max(max_row_seen, cell_row)
                    if cell.value is None:
                        continue
                    raw = str(cell.value).strip()
                    if not raw:
                        continue
                    cells.append((cell.coordinate, cell_row, raw))
                    if cell_row is not None:
                        non_empty_rows.add(cell_row)

            if cells:
                cells_per_row: dict[int, int] = {}
                for _coord, row_idx, _text in cells:
                    cells_per_row[row_idx] = cells_per_row.get(row_idx, 0) + 1
                out.append(
                    _build_overview_fragment(
                        sheet_name=sheet_name,
                        non_empty_rows=sorted(non_empty_rows),
                        max_row_seen=max_row_seen,
                        cells_per_row=cells_per_row,
                    )
                )

            for coord, _row, text in cells:
                out.append(
                    TextFragment(
                        text=text,
                        location_hint=f"{sheet_name}!{coord}",
                    )
                )
    finally:
        wb.close()
    return out


def _build_overview_fragment(
    sheet_name: str,
    non_empty_rows: list[int],
    max_row_seen: int,
    cells_per_row: dict[int, int],
) -> TextFragment:
    """Render a structural summary the LLM can use to spot blank-row templates.

    The text always lists non-empty row numbers; if the sheet looks like
    "small header + lots of blank rows" we tag it explicitly so the LLM's
    Mode B branch fires reliably (without this hint the model only ever
    sees the header cells and can't distinguish a tiny sheet from a
    blank-row entry form).

    Two ways the "header + blank" pattern shows up in real templates:

    1. Non-empty rows ≪ used rows — user once typed/cleared something
       further down, so openpyxl's max_row reports a bigger number than
       the rows actually carrying data right now.
    2. Only one non-empty row exists at all — fresh template, openpyxl
       didn't pad past the header. ``max_row_seen == 1`` but the LLM still
       needs to know that "data is meant to go in row 2".

    Both cases land on the same hint string: "rows N have content; data
    should be filled starting from row N+1".
    """
    parts = [f'工作表 "{sheet_name}"', f"已使用至第 {max_row_seen} 行"]
    if non_empty_rows:
        rows_repr = ",".join(str(r) for r in non_empty_rows[:8])
        if len(non_empty_rows) > 8:
            rows_repr += f",…(共 {len(non_empty_rows)} 行)"
        parts.append(f"有内容的行号: [{rows_repr}]")

    if not non_empty_rows:
        return TextFragment(text="; ".join(parts), location_hint=f"{sheet_name}!__overview__")

    blank_rows = max_row_seen - len(non_empty_rows)
    looks_like_header_form = False
    if blank_rows >= 1 and len(non_empty_rows) <= 3:
        # Pattern 1: blank gaps detected by openpyxl.
        looks_like_header_form = True
    elif len(non_empty_rows) == 1 and cells_per_row.get(non_empty_rows[0], 0) >= 2:
        # Pattern 2: single row with ≥2 cells — typical "label row only" form.
        looks_like_header_form = True

    if looks_like_header_form:
        first_blank = next(
            (r for r in range(1, max(max_row_seen, non_empty_rows[-1]) + 2) if r not in non_empty_rows),
            non_empty_rows[-1] + 1,
        )
        rows_str = ",".join(str(r) for r in non_empty_rows)
        hint = f"⚠️ 结构提示: 第 {rows_str} 行有内容,其余行均为空白。 这通常是『列头 + 空白填充行』表单结构,数据应填入第 {first_blank} 行起。"
        parts.append(hint)

    return TextFragment(text="; ".join(parts), location_hint=f"{sheet_name}!__overview__")


def chunk_for_prompt(fragments: list[TextFragment], max_chars: int = 8000) -> Iterator[list[TextFragment]]:
    """Split fragments into batches the LLM can comfortably ingest.

    `max_chars` is a soft budget on the joined "text + hint" payload per
    batch. Big templates (1000+ rows) would otherwise blow the context
    window — chunking keeps each LLM call bounded and lets the caller
    merge the field sets with simple deduplication on `original_text`.
    """
    batch: list[TextFragment] = []
    used = 0
    for frag in fragments:
        cost = len(frag.text) + len(frag.location_hint) + 4  # tab/newline overhead
        if batch and used + cost > max_chars:
            yield batch
            batch, used = [], 0
        batch.append(frag)
        used += cost
    if batch:
        yield batch
